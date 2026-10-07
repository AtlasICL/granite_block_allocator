"""Reading block lists from CSV and Excel files.

The loader is deliberately forgiving about how a file was produced (Excel
exports, different locales, stray blank rows) and strict about the data
itself: every block needs an identifier and a positive weight. When something
is wrong, the error message says what was found and where, so the person
using the app can fix the file without help.
"""

from __future__ import annotations

import csv
import io
import math
import re
from collections import Counter
from collections.abc import Sequence
from pathlib import Path

Block = tuple[str, float]

CSV_EXTENSIONS = (".csv", ".txt", ".tsv")
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
SUPPORTED_EXTENSIONS = CSV_EXTENSIONS + EXCEL_EXTENSIONS

# Header names accepted for each required column, after normalisation
# (lower case, no spaces/underscores/hyphens/dots, unit suffixes removed).
BLOCK_ID_NAMES = ("blockno", "blocknumber", "blocknum", "blockid")
WEIGHT_NAMES = ("weight", "wt")

# How many individual problems to list before summarising the rest.
_MAX_REPORTED_PROBLEMS = 5


class BlockFileError(ValueError):
    """The block file could not be read. The message is meant for the user."""


# --------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------


def load_blocks(path: str | Path) -> list[Block]:
    """Load a block list and return it as ``(BlockNo, Weight)`` tuples.

    ``BlockNo`` is kept as text so identifiers such as ``0612`` or ``A-17``
    survive unchanged. Fully blank rows are skipped.

    Raises:
        BlockFileError: if the file can't be read, the required columns can't
            be found, or any row has a missing or invalid value.
    """
    path = Path(path)
    if not path.exists():
        raise BlockFileError(f"The file '{path.name}' could not be found.")

    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise BlockFileError(
            "Old-style Excel files (.xls) can't be read. In Excel, use "
            "File > Save As and choose 'Excel Workbook (.xlsx)' or 'CSV'."
        )
    if suffix in EXCEL_EXTENSIONS:
        rows = _read_excel_rows(path)
    else:
        rows = _read_csv_rows(path)

    return _blocks_from_rows(rows)


def find_duplicate_ids(blocks: Sequence[Block]) -> list[str]:
    """Return the block numbers that appear more than once, in file order."""
    counts = Counter(block_id for block_id, _ in blocks)
    return list(dict.fromkeys(b for b, _ in blocks if counts[b] > 1))


def normalise_header(name: str) -> str:
    """Normalise a column header for matching.

    ``" Block No "``, ``"block_no"`` and ``"BLOCK-NO"`` all become
    ``"blockno"``; a trailing unit such as ``"Weight (t)"`` or
    ``"Weight [kg]"`` is dropped.
    """
    text = name.strip().lower()
    text = re.sub(r"\s*[\(\[].*?[\)\]]\s*$", "", text)
    return re.sub(r"[\s_\-.]+", "", text)


def parse_weight(text: str) -> float | None:
    """Parse a weight written with either ``.`` or ``,`` as the decimal mark.

    Returns ``None`` if the text isn't a plain number. Thousands separators
    are accepted when both marks are present (``1.234,5`` or ``1,234.5``).
    """
    s = text.strip().replace(" ", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        # Whichever mark comes last is the decimal mark.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        if s.count(",") > 1:
            return None
        s = s.replace(",", ".")
    if not re.fullmatch(r"[+-]?(\d+(\.\d*)?|\.\d+)", s):
        return None
    return float(s)


# --------------------------------------------------------------------------
# Reading raw rows
# --------------------------------------------------------------------------

# Each raw row is (row number as shown in Excel / a text editor, cell values).
_RawRow = tuple[int, list[str]]
_RawRows = tuple[list[str], list[_RawRow]]


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel's "CSV (Comma delimited)" uses the Windows code page. Only the
        # BlockNo and Weight columns matter, so a lossy decode is acceptable.
        return data.decode("cp1252", errors="replace")


def _detect_delimiter(header_line: str) -> str:
    """Pick the delimiter used in the header line: ``,`` ``;`` or tab.

    Excel in many European locales saves "CSV" files with semicolons.
    """
    counts = {d: header_line.count(d) for d in (",", ";", "\t")}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] > 0 else ","


def _read_csv_rows(path: Path) -> _RawRows:
    try:
        text = _decode(path.read_bytes())
    except OSError as exc:
        raise BlockFileError(f"The file '{path.name}' could not be opened: {exc.strerror}.") from exc

    lines = text.splitlines()
    first = next((line for line in lines if line.strip()), None)
    if first is None:
        return [], []

    reader = csv.reader(io.StringIO(text), delimiter=_detect_delimiter(first))
    header: list[str] | None = None
    rows: list[_RawRow] = []
    for cells in reader:
        if not any(c.strip() for c in cells):
            continue
        if header is None:
            header = cells
        else:
            rows.append((reader.line_num, cells))
    return header or [], rows


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def _read_excel_rows(path: Path) -> _RawRows:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise BlockFileError("Excel support needs the 'openpyxl' package.") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different types
        raise BlockFileError(
            f"'{path.name}' could not be read as an Excel workbook. "
            "Check that it opens in Excel, or save it as CSV."
        ) from exc

    try:
        sheet = workbook.active
        if sheet is None:
            return [], []
        header: list[str] | None = None
        rows: list[_RawRow] = []
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [_cell_to_text(v) for v in values]
            if not any(c.strip() for c in cells):
                continue
            if header is None:
                header = cells
            else:
                rows.append((row_number, cells))
        return header or [], rows
    finally:
        workbook.close()


# --------------------------------------------------------------------------
# Turning rows into blocks
# --------------------------------------------------------------------------


def _find_column(header: Sequence[str], accepted: Sequence[str], label: str) -> int:
    matches = [i for i, name in enumerate(header) if normalise_header(name) in accepted]
    if len(matches) == 1:
        return matches[0]

    shown = ", ".join(f"'{h.strip()}'" for h in header if h.strip()) or "none"
    if not matches:
        raise BlockFileError(
            f"Couldn't find a {label} column. Columns in this file: {shown}.\n\n"
            f"Rename the right column to '{label}' (capitals and spaces don't matter)."
        )
    names = " and ".join(f"'{header[i].strip()}'" for i in matches)
    raise BlockFileError(
        f"More than one column could be the {label} column ({names}). Rename or remove one of them."
    )


def _summarise(problems: list[str]) -> str:
    shown = problems[:_MAX_REPORTED_PROBLEMS]
    extra = len(problems) - len(shown)
    text = "\n".join(f"  • {p}" for p in shown)
    if extra:
        text += f"\n  • …and {extra} more"
    return text


def _blocks_from_rows(raw: _RawRows) -> list[Block]:
    header, rows = raw
    if not header:
        return []

    id_col = _find_column(header, BLOCK_ID_NAMES, "BlockNo")
    weight_col = _find_column(header, WEIGHT_NAMES, "Weight")

    blocks: list[Block] = []
    problems: list[str] = []
    for row_number, cells in rows:
        block_id = cells[id_col].strip() if id_col < len(cells) else ""
        weight_text = cells[weight_col].strip() if weight_col < len(cells) else ""

        if not block_id:
            problems.append(f"Row {row_number}: BlockNo is empty")
            continue
        if not weight_text:
            problems.append(f"Row {row_number} (block {block_id}): Weight is empty")
            continue

        weight = parse_weight(weight_text)
        if weight is None or not math.isfinite(weight):
            problems.append(f"Row {row_number} (block {block_id}): '{weight_text}' is not a number")
            continue
        if weight <= 0:
            problems.append(f"Row {row_number} (block {block_id}): Weight must be more than 0")
            continue

        blocks.append((block_id, weight))

    if problems:
        noun = "row has a problem" if len(problems) == 1 else "rows have problems"
        raise BlockFileError(f"{len(problems)} {noun}:\n{_summarise(problems)}")
    return blocks
